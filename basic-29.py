# 51. Automation / Scripting
# Automation means writing Python programs to perform repetitive tasks automatically instead of doing them manually.

# 1. File Automation
# Automating file creation, reading, updating, deleting, and renaming.

# Create a File
with open("demo.txt", "w") as file:
    file.write("Hello World")
# Read a File
with open("demo.txt", "r") as file:
    print(file.read())
# Rename a File
import os

os.rename("demo.txt", "new_demo.txt")
# Delete a File
import os

os.remove("new_demo.txt")

# 2. Folder Automation
# - Creating, moving, deleting, and organizing folders.

# Create Folder
import os

os.mkdir("Reports")

# Create Nested Folders
os.makedirs("2026/October/Reports")

# Delete Empty Folder
os.rmdir("Reports")

# Delete Folder with Contents
import shutil

shutil.rmtree("2026")

# List Folder Contents
for item in os.listdir():
    print(item)

# 3. OS Commands
# - Python can execute operating system commands.
# Windows
import os
os.system("dir")

# Linux
import os
os.system("ls")

# 4. os Module
# - Provides interaction with the operating system.
# Current Working Directory
import os
print(os.getcwd())

# Change Directory
os.chdir("D:/Projects")

# Environment Variables
import os
print(os.environ["PATH"])

# Check File Exists
print(os.path.exists("demo.txt"))

# 5. pathlib
# - Modern, object-oriented path handling.
# Create Path
from pathlib import Path

path = Path("data/file.txt")
print(path)

# Check File Exists
from pathlib import Path

file = Path("data.txt")

print(file.exists())

# File Extension
print(file.suffix)

# File Name
print(file.name)

# Parent Folder
print(file.parent)

# Create Directory
# Path("logs").mkdir(exist_ok=True)

# 6. shutil
# - Used for high-level file operations.
# Copy File
import shutil
shutil.copy("a.txt", "b.txt")

# Move File
shutil.move("a.txt", "Backup/a.txt")

# Copy Entire Directory
shutil.copytree("Source", "Backup")

# Delete Directory
shutil.rmtree("Backup")

# 7. subprocess
# - Runs external applications and system commands safely.
# Execute Command
# import subprocess
# subprocess.run(["ipconfig"])

# Capture Output
result = subprocess.run(
    ["python", "--version"],
    capture_output=True,
    text=True
)

print(result.stdout)

# Open Notepad (Windows)
subprocess.run(["notepad.exe"])

# Open Chrome
subprocess.run([
    r"C:\Program Files\Google\Chrome\Application\chrome.exe"
])

# 8. Excel Automation
# - Using openpyxl.
# Install: pip install openpyxl

# Create Excel File
from openpyxl import Workbook

wb = Workbook()
ws = wb.active

ws["A1"] = "Name"
ws["B1"] = "Marks"

wb.save("students.xlsx")

# Read Excel
from openpyxl import load_workbook

wb = load_workbook("students.xlsx")
ws = wb.active

print(ws["A1"].value)

# Add Data
ws.append(["Aman", 90])
ws.append(["Rahul", 85])

wb.save("students.xlsx")

# 9. PDF Automation
# - Using PyPDF2.
# Install: pip install PyPDF2

# Read PDF
from PyPDF2 import PdfReader

reader = PdfReader("sample.pdf")

for page in reader.pages:
    print(page.extract_text())

# Merge PDFs
from PyPDF2 import PdfMerger

merger = PdfMerger()

merger.append("file1.pdf")
merger.append("file2.pdf")

merger.write("merged.pdf")
merger.close()

# 10. Email Automation
# - Using smtplib.
# Send Email
import smtplib

server = smtplib.SMTP("smtp.gmail.com", 587)

server.starttls()

server.login(
    "your_email@gmail.com",
    "app_password"
)

server.sendmail(
    "your_email@gmail.com",
    "receiver@gmail.com",
    "Hello from Python"
)

server.quit()

# Using EmailMessage
from email.message import EmailMessage

msg = EmailMessage()

msg["Subject"] = "Test"
msg["From"] = "sender@gmail.com"
msg["To"] = "receiver@gmail.com"

msg.set_content("Hello")

# 11. Web Automation
# - Using requests and BeautifulSoup.
# Install: pip install requests beautifulsoup4
# Download Web Page
import requests

response = requests.get(
    "https://example.com"
)

print(response.text)

# Parse HTML
from bs4 import BeautifulSoup

html = "<h1>Hello</h1>"

soup = BeautifulSoup(
    html,
    "html.parser"
)

print(soup.h1.text)

# Scrape Website Title
import requests
from bs4 import BeautifulSoup

response = requests.get(
    "https://example.com"
)

soup = BeautifulSoup(
    response.text,
    "html.parser"
)

print(soup.title.text)

# Browser Automation (Selenium)
# pip install selenium
from selenium import webdriver

driver = webdriver.Chrome()

driver.get(
    "https://google.com"
)

print(driver.title)

driver.quit()

# Real-World Automation Projects
# Beginner
# Auto file organizer
# Bulk file renamer
# Folder backup script
# Daily log generator
# Excel marks calculator

# Intermediate
# PDF merger tool
# Email sender
# Website scraper
# Expense tracker with Excel
# System monitoring script

# Advanced
# WhatsApp automation
# LinkedIn job scraper
# Automated report generator
# Email attachment downloader
# Scheduled backup system

# Interview Questions
# Q1. Difference between os and pathlib?
# os → Traditional approach
# pathlib → Modern object-oriented approach
# os
os.path.exists("data.txt")
# pathlib
Path("data.txt").exists()

# Q2. Why use subprocess instead of os.system()?
# More secure
# Captures output
# Better error handling

# Q3. What is shutil used for?
# High-level file and folder operations:
# Copy
# Move
# Archive
# Delete

# Q4. What is automation scripting?
# - Writing programs that perform repetitive tasks automatically.

# Q5. Which library is commonly used for Excel automation?
# openpyxl
# pandas

# --------------------------------------------------- #
# 52. Web Scraping
# Web Scraping is the process of automatically extracting data from websites using Python programs.

# 1. HTML (HyperText Markup Language)
# - HTML is the structure of a web page.
# <html>
# <body>
#     <h1>My Website</h1>
#     <p>Hello World</p>
# </body>
# </html>
# Python can read this HTML and extract required information.

# 2. DOM (Document Object Model)
# - DOM is a tree representation of HTML.
# <html>
#     <body>
#         <h1>Title</h1>
#         <p>Paragraph</p>
#     </body>
# </html>

# DOM Tree:
# html
#  └── body
#       ├── h1
#       └── p
# Scraping libraries navigate this tree to find elements.

# 3. Requests
# - Used to download webpage content.
# Installation: pip install requests
import requests

url = "https://example.com"

response = requests.get(url)

print(response.status_code)
print(response.text)

# Common Methods:
requests.get()
requests.post()
requests.put()
requests.delete()

# Useful Attributes:
response.text
response.content
response.status_code
response.headers

# 4. BeautifulSoup
# - Used for parsing HTML.
# Installation: pip install beautifulsoup4
from bs4 import BeautifulSoup
import requests

url = "https://example.com"

html = requests.get(url).text

soup = BeautifulSoup(html, "html.parser")

print(soup.title.text)

# Finding Elements:
soup.find("h1")

soup.find_all("p")

soup.find(id="main")

soup.find(class_="product")

# 5. CSS Selectors
# - CSS selectors are used to target HTML elements.
# HTML:
# <div class="product">
#     <h2>Laptop</h2>
# </div>

# Python: soup.select(".product")
soup.select("h1")

soup.select(".product")

soup.select("#main")

soup.select("div h2")

# Selector Examples:
# .product
# #main
# div h2
# ul li
# a[href]

# 6. XPath
# - Alternative way to locate elements.
# Installation: pip install lxml
from lxml import html
import requests

page = requests.get("https://example.com")

tree = html.fromstring(page.content)

title = tree.xpath("//h1/text()")

print(title)

# Common XPath:
# //h1
#
# //div[@class='product']
#
# //a/@href
#
# //input[@id='email']
#
# //table/tr

# 7. Pagination
# - Many websites split data across pages.
# Example URLs:
page=1
page=2
page=3

# Scraping Multiple Pages:
import requests
from bs4 import BeautifulSoup

for page in range(1, 6):
    url = f"https://site.com/products?page={page}"

    html = requests.get(url).text

    soup = BeautifulSoup(html, "html.parser")

    products = soup.find_all("h2")

    for product in products:
        print(product.text)

# 8. Selenium
# Selenium automates a real browser.
# Used when:
# JavaScript loads content
# Login required
# Buttons must be clicked
# Installation: pip install selenium
from selenium import webdriver

driver = webdriver.Chrome()

driver.get("https://example.com")

print(driver.title)

driver.quit()
# Finding Elements:
from selenium.webdriver.common.by import By

driver.find_element(By.ID, "email")

driver.find_element(By.CLASS_NAME, "product")

driver.find_element(By.XPATH, "//button")
# Clicking Buttons: button.click()
# Typing: search.send_keys("Python")

# 9. Playwright
# Modern alternative to Selenium.
# Advantages:
# - Faster
# - More stable
# - Supports multiple browsers
# - Excellent for dynamic websites
# Installation:
# pip install playwright
# playwright install
from playwright.sync_api import sync_playwright

with sync_playwright() as p:

    browser = p.chromium.launch()

    page = browser.new_page()

    page.goto("https://example.com")

    print(page.title())

    browser.close()
# Extract Data:
title = page.locator("h1").text_content()
print(title)

# 10. Handling Dynamic Websites
# - Dynamic websites load content using JavaScript.
# Problem:
requests.get()
# may return: <div id="products"></div>
# because the data is loaded later by JavaScript.
# Solutions:
# Use Selenium
driver.get(url)
# Wait for content:
from selenium.webdriver.support.ui import WebDriverWait

WebDriverWait(driver, 10)

# Use Playwright
page.goto(url)

page.wait_for_load_state("networkidle")

# Inspect API Calls
# Open: F12 → Network → Fetch/XHR
# Find API: https://api.site.com/products
# Then directly request:
import requests

data = requests.get(api_url).json()

print(data)

# This is usually the fastest approach.

# Complete Example: Requests + BeautifulSoup
import requests
from bs4 import BeautifulSoup

url = "https://quotes.toscrape.com"

response = requests.get(url)

soup = BeautifulSoup(response.text, "html.parser")

quotes = soup.select(".quote")

for quote in quotes:
    text = quote.select_one(".text").text
    author = quote.select_one(".author").text

    print(author, ":", text)

# Interview Questions
# Q1. Difference between Requests and BeautifulSoup?
# Requests: Downloads webpage.
# BeautifulSoup: Parses and extracts data from HTML.

# Q2. When should Selenium be used?
# When content is loaded using JavaScript or browser interaction is needed.

# Q3. Selenium vs Playwright?
# Selenium - Playwright
# Older - Newer
# Slower - Faster
# More setup - Easier
# Widely adopted - Growing rapidly

# Q4. XPath vs CSS Selectors?
# XPath is more powerful.
# CSS selectors are easier to read.
# Both are used for locating elements.

# Q5. Why is scraping sometimes difficult?
# Because websites may use:
# JavaScript rendering
# Captchas
# Rate limiting
# Authentication
# Infinite scrolling

# Mini Project Ideas
# News Headline Scraper
# Amazon Product Tracker
# Weather Data Scraper
# Job Listing Scraper
# Cryptocurrency Price Monitor
# LinkedIn Job Scraper
# E-commerce Product Comparison Tool
# Stock Market Data Collector